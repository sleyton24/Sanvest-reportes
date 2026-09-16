"""Preventas SV155 salen de Estadística.Pagado; ML/SV99 siguen congeladas.

Correr desde la raíz del repo:
    python -m unittest tests.test_dv_preventas
"""
from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd
from sqlalchemy import create_engine

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from etl import informes_dv as I  # noqa: E402
from etl.connect_dv import apply_dv  # noqa: E402

SV155 = "Sta. Victoria 155"
ML = "Millalongo"
SV99 = "Sta. Victoria 99"
PAGADO_FILE = 98765.43
PAGADO_PREV = 11111.0
LINEA_SV155_PREV = 226074.0
EGRESOS_PREV = 500000.0


def _fill_estad_sheet(ws, *, pagado: float | None, pagado_col: int = 16) -> None:
    ws.cell(1, 1, "Deptos.")
    ws.cell(1, 2, 154)
    ws.cell(1, 4, 42)
    ws.cell(1, 6, 6)
    ws.cell(1, 8, 106)
    ws.cell(1, 16, "Pagado")
    ws.cell(2, 11, "Sub-Total Precio Venta")
    ws.cell(2, 14, 210000)
    ws.cell(3, 11, "Falta x Vender")
    ws.cell(3, 14, 40000)
    ws.cell(3, pagado_col, 8000)
    ws.cell(4, 1, "Totales")
    ws.cell(4, pagado_col, pagado)


def _write_estadistica(path: Path, *, sheet: str = "SV.155", pagado: float = PAGADO_FILE,
                       pagado_col: int = 16) -> None:
    """Estadística mínima: TOTALES col Pagado (P=16 por defecto) + Deptos.

    `sheet` es la hoja de SV155 (puede ser alias Sv155 / SV155); no se crea
    también SV.155, para que el resolver de alias no tome una hoja vacía.
    """
    wb = openpyxl.Workbook()
    names = ["MI.72", "SV.99", sheet]
    first = True
    for name in names:
        ws = wb.active if first else wb.create_sheet(name)
        first = False
        ws.title = name
        val = pagado if name == sheet else 100.0
        _fill_estad_sheet(ws, pagado=val, pagado_col=pagado_col)
    wb.save(path)
    wb.close()


def _write_rentab(path: Path, year: int = 2026, month: int = 8) -> None:
    wb = openpyxl.Workbook()
    first = True
    for name in ("MI.72", "SV.99", "SV.155"):
        ws = wb.active if first else wb.create_sheet(name)
        first = False
        ws.title = name
        ws.cell(5, 1, date(year, month, 1))
        ws.cell(5, 4, 10.0)   # Socio
        ws.cell(5, 6, 20.0)   # Danacorp
        ws.cell(5, 8, 5.0)    # Ventas
    wb.save(path)
    wb.close()


def _uf_row(proj: str, fid: int, sub: str, monto: float) -> dict:
    y, m = divmod(fid, 100)
    return {
        "Mes de carga": m, "Año de carga": y, "Tipo de datos": "REAL",
        "Nombre proyecto": proj, "Fecha": f"{y}-{m:02d}-01", "Fecha ID": fid,
        "Categoria": "FONDOS" if sub != "EGRESOS A LA FECHA" else "USOS",
        "SUBCATEGORIA": sub, "Monto": monto,
    }


class EstadisticaPagado(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_lee_pagado_col_p_totales(self):
        p = self.dir / "Estadística de Ventas.xlsx"
        _write_estadistica(p)
        est = I.estadistica_venta(p, "SV.155")
        self.assertAlmostEqual(est["pagado"], PAGADO_FILE)
        self.assertAlmostEqual(est["recaud"], PAGADO_FILE)

    def test_alias_hoja_sv155(self):
        p = self.dir / "estadistica.xlsx"
        _write_estadistica(p, sheet="Sv155")
        # PROJECTS pide SV.155; el alias sin punto debe resolver
        est = I.estadistica_venta(p, "SV.155")
        self.assertAlmostEqual(est["pagado"], PAGADO_FILE)

    def test_pagado_por_encabezado_si_no_esta_en_p(self):
        p = self.dir / "estad.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "SV.155"
        ws.cell(1, 12, "Pagado")          # col L
        ws.cell(1, 1, "Deptos.")
        ws.cell(1, 4, 10)
        ws.cell(2, 1, "Totales")
        ws.cell(2, 12, 333.5)
        wb.save(p)
        wb.close()
        est = I.estadistica_venta(p, "SV.155")
        self.assertAlmostEqual(est["pagado"], 333.5)

    def test_preventas_sv155_usa_archivo_no_fallback(self):
        p = self.dir / "Estadística de Ventas.xlsx"
        _write_estadistica(p)
        val, origen = I.preventas_sv155(p, PAGADO_PREV)
        self.assertAlmostEqual(val, PAGADO_FILE)
        self.assertEqual(origen, "estadistica.pagado")

    def test_preventas_sv155_sin_archivo_lleva_fallback(self):
        val, origen = I.preventas_sv155(None, PAGADO_PREV)
        self.assertAlmostEqual(val, PAGADO_PREV)
        self.assertEqual(origen, "carry_forward")


class UsosYFondos(unittest.TestCase):
    flujo = {"socio": 10.0, "danacorp": 20.0, "ventas": 0.0}

    def _by_sub(self, rows):
        return {r["SUBCATEGORIA"]: r["Monto"] for r in rows}

    def test_sv155_preventas_del_archivo_recalcula_capital(self):
        rows = I.usos_y_fondos_rows(
            {"egresos": EGRESOS_PREV}, SV155, 2026, 8, self.flujo,
            PAGADO_FILE, LINEA_SV155_PREV)
        by = self._by_sub(rows)
        egresos = EGRESOS_PREV + 30.0
        self.assertAlmostEqual(by["PREVENTAS"], PAGADO_FILE)
        self.assertAlmostEqual(by["LÍNEA DE CRÉDITO GIRADA"], LINEA_SV155_PREV)
        self.assertAlmostEqual(by["CAPITAL SOCIOS FONDOS"],
                               egresos - LINEA_SV155_PREV - PAGADO_FILE)

    def test_ml_y_sv99_siguen_congeladas(self):
        ml = self._by_sub(I.usos_y_fondos_rows(
            {"egresos": 1.0}, ML, 2026, 8, self.flujo, PAGADO_FILE, None, 225597.0))
        sv = self._by_sub(I.usos_y_fondos_rows(
            {"egresos": 1.0}, SV99, 2026, 8, self.flujo, PAGADO_FILE, None))
        self.assertAlmostEqual(ml["PREVENTAS"], I.PREVENTAS_FROZEN[ML])
        self.assertAlmostEqual(sv["PREVENTAS"], I.PREVENTAS_FROZEN[SV99])
        self.assertAlmostEqual(sv["LÍNEA DE CRÉDITO GIRADA"], I.LINEA_FROZEN[SV99])


class ApplyDvPreventas(unittest.TestCase):
    """Smoke: apply_dv escribe PREVENTAS SV155 desde el archivo, no el mes anterior."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.engine = create_engine("sqlite:///:memory:")
        prev_fid = 202607
        rows = []
        for proj, pag, lin, egr in (
            (ML, I.PREVENTAS_FROZEN[ML], 225597.0, 400000.0),
            (SV99, I.PREVENTAS_FROZEN[SV99], I.LINEA_FROZEN[SV99], 450000.0),
            (SV155, PAGADO_PREV, LINEA_SV155_PREV, EGRESOS_PREV),
        ):
            rows += [
                _uf_row(proj, prev_fid, "EGRESOS A LA FECHA", egr),
                _uf_row(proj, prev_fid, "LÍNEA DE CRÉDITO GIRADA", lin),
                _uf_row(proj, prev_fid, "PREVENTAS", pag),
                _uf_row(proj, prev_fid, "CAPITAL SOCIOS FONDOS", egr - lin - pag),
            ]
        pd.DataFrame(rows).to_sql("dv_uso_y_fondo", self.engine, index=False)
        for table, cols in (
            ("dv_evolucion_de_costos",
             ["Mes de carga", "Año de carga", "Tipo de datos", "Nombre proyecto",
              "Año", "mes", "Periodo", "PPTO_DE_COSTOS", "PROYECCIÓN_DE_COSTOS",
              "COSTOS_REALES", "Fecha Id"]),
            ("dv_escrituras",
             ["Mes de carga", "Año de carga", "Tipo de datos", "Nombre proyecto",
              "Versión", "VENTAS_ACUMULADAS", "UF_RECAUDADAS         ",
              "UF_POR_RECAUDAR", "PROYECCIÓN_VENTA_TOTAL(UF)",
              "UNIDADES_ESCRITURADAS_RECAUDADAS", "UNIDADES_ESCRITURADAS_FIRMADAS",
              "RESERVAS_Y_PROMESAS", "Periodo", "Fehca ID "]),
            ("dv_ventas",
             ["Mes de carga", "Año de carga", "Tipo de datos", "Nombre proyecto",
              "Versión", "VENTAS_ACUMULADAS", "UF_RECAUDADAS         ",
              "UF_POR_RECAUDAR", "PROYECCIÓN_VENTA_TOTAL(UF)",
              "UNIDADES_ESCRITURADAS_RECAUDADAS", "UNIDADES_ESCRITURADAS_FIRMADAS",
              "RESERVAS_Y_PROMESAS", "Año Mes ", "Fecha ID "]),
            ("dv_kpis",
             ["Mes", "Año", "Tipo de datos", "Nombre proyecto", "Versión",
              "VENTAS NETAS_DEL_MES", "UNIDADES_VENDIDAS", "AVANCE_VENTAS_(UNIDADES)%",
              "UF/M2_VENTA", "UNIDADES TOTALES", "Periodo", "Fecha ID"]),
            ("dv_construccion",
             ["Nombre proyecto", "Fecha ID ", "Periodo", "Mes de carga ", "Año de carga"]),
            ("dv_indicadores_financieros",
             ["Nombre proyecto", "fecha ID ", "Periodo", "Mes", "Año"]),
        ):
            pd.DataFrame(columns=cols).to_sql(table, self.engine, index=False)

    def tearDown(self):
        self.tmp.cleanup()
        self.engine.dispose()

    def test_apply_dv_sv155_preventas_desde_estadistica(self):
        rent = self.dir / "Rentabilidad Inversiones Proyectos.xlsx"
        estad = self.dir / "Estadística de Ventas.xlsx"
        _write_rentab(rent)
        _write_estadistica(estad)
        out = apply_dv(self.engine, {"rentabilidad": rent, "estadistica": estad,
                                     "escrituracion": {}})
        self.assertEqual(out["periodo"], 202608)
        self.assertEqual(out["sv155_preventas"]["origen"], "estadistica.pagado")
        self.assertAlmostEqual(out["sv155_preventas"]["valor"], PAGADO_FILE)

        uf = pd.read_sql_query('SELECT * FROM "dv_uso_y_fondo"', self.engine)
        cur = uf[uf["Fecha ID"] == 202608]

        def monto(proj, sub):
            r = cur[(cur["Nombre proyecto"] == proj) & (cur["SUBCATEGORIA"] == sub)]
            return float(r["Monto"].iloc[0])

        self.assertAlmostEqual(monto(SV155, "PREVENTAS"), PAGADO_FILE)
        self.assertAlmostEqual(monto(SV155, "LÍNEA DE CRÉDITO GIRADA"), LINEA_SV155_PREV)
        self.assertAlmostEqual(monto(ML, "PREVENTAS"), I.PREVENTAS_FROZEN[ML])
        self.assertAlmostEqual(monto(SV99, "PREVENTAS"), I.PREVENTAS_FROZEN[SV99])
        egresos = EGRESOS_PREV + 30.0
        self.assertAlmostEqual(monto(SV155, "CAPITAL SOCIOS FONDOS"),
                               egresos - LINEA_SV155_PREV - PAGADO_FILE)

    def test_apply_dv_sin_estadistica_arrastra_preventas(self):
        rent = self.dir / "Rentabilidad Inversiones Proyectos.xlsx"
        _write_rentab(rent)
        out = apply_dv(self.engine, {"rentabilidad": rent, "escrituracion": {}})
        self.assertEqual(out["sv155_preventas"]["origen"], "carry_forward")
        self.assertAlmostEqual(out["sv155_preventas"]["valor"], PAGADO_PREV)
        uf = pd.read_sql_query(
            'SELECT "Monto" FROM "dv_uso_y_fondo" WHERE "Nombre proyecto"=:p '
            'AND "Fecha ID"=202608 AND "SUBCATEGORIA"=\'PREVENTAS\'',
            self.engine, params={"p": SV155})
        self.assertAlmostEqual(float(uf["Monto"].iloc[0]), PAGADO_PREV)


if __name__ == "__main__":
    unittest.main()
